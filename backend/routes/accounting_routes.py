from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import StreamingResponse
from datetime import datetime, timezone
from typing import Optional
import logging

from database import db
from auth import get_current_admin
from models import CashWithdrawalCreate
import uuid

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api")

@router.get("/accounting/report")
async def get_accounting_report(
    gym_id: Optional[str] = None, date_from: Optional[str] = None,
    date_to: Optional[str] = None, category: Optional[str] = None,
    admin: dict = Depends(get_current_admin)
):
    query = {"payment_status": "paid"}
    if admin["role"] != "super_admin":
        query["gym_id"] = admin.get("gym_id")
    elif gym_id:
        query["gym_id"] = gym_id
    if date_from:
        query["created_at"] = {"$gte": date_from}
    if date_to:
        if "created_at" in query:
            query["created_at"]["$lte"] = date_to + "T23:59:59"
        else:
            query["created_at"] = {"$lte": date_to + "T23:59:59"}
    transactions = await db.payment_transactions.find(query, {"_id": 0}).sort("created_at", -1).to_list(5000)

    # Build plan_id -> activity map. "Activity" is the plan's explicit category when the admin set a custom one,
    # otherwise it falls back to the plan's own name. This way La Fabrika's existing plans
    # ("Kickboxing", "Boxeo", "Capoeira"...) appear automatically as activities with zero config.
    plan_query = {}
    tgt_gym_id = admin.get("gym_id") if admin["role"] != "super_admin" else gym_id
    if tgt_gym_id:
        plan_query["gym_id"] = tgt_gym_id
    plans_cat = await db.plans.find(plan_query, {"_id": 0, "id": 1, "name": 1, "category": 1}).to_list(500)
    plan_cat_map = {}
    for p in plans_cat:
        cat = p.get("category")
        # Treat the default/unset "General" as not-set and fall back to the plan name
        if not cat or cat.strip().lower() == "general":
            cat = p.get("name") or "Sin plan"
        plan_cat_map[p["id"]] = cat

    # Attach activity to each transaction. Fall back to transaction's own plan_name
    # so legacy or deleted plans still show up as a grouping instead of "General".
    for t in transactions:
        activity = plan_cat_map.get(t.get("plan_id"))
        if not activity:
            activity = t.get("plan_name") or "Sin plan"
        t["category"] = activity

    # Optional filter by activity
    if category and category != "all":
        transactions = [t for t in transactions if (t.get("category") or "Sin plan") == category]
    
    # Include POS sales
    pos_query = {}
    if admin["role"] != "super_admin":
        pos_query["gym_id"] = admin.get("gym_id")
    elif gym_id:
        pos_query["gym_id"] = gym_id
    if date_from:
        pos_query["created_at"] = {"$gte": date_from}
    if date_to:
        if "created_at" in pos_query:
            pos_query["created_at"]["$lte"] = date_to + "T23:59:59"
        else:
            pos_query["created_at"] = {"$lte": date_to + "T23:59:59"}
    pos_sales = await db.pos_sales.find(pos_query, {"_id": 0}).sort("created_at", -1).to_list(5000)
    
    # Cash withdrawals
    withdrawal_query = {}
    if admin["role"] != "super_admin":
        withdrawal_query["gym_id"] = admin.get("gym_id")
    elif gym_id:
        withdrawal_query["gym_id"] = gym_id
    if date_from:
        withdrawal_query["created_at"] = {"$gte": date_from}
    if date_to:
        if "created_at" in withdrawal_query:
            withdrawal_query["created_at"]["$lte"] = date_to + "T23:59:59"
        else:
            withdrawal_query["created_at"] = {"$lte": date_to + "T23:59:59"}
    withdrawals = await db.cash_withdrawals.find(withdrawal_query, {"_id": 0}).sort("created_at", -1).to_list(1000)
    
    total_revenue = sum(t.get("amount", 0) for t in transactions)
    cash_total = sum(t.get("amount", 0) for t in transactions if t.get("payment_method") == "cash")
    card_total = sum(t.get("amount", 0) for t in transactions if t.get("payment_method") == "card_reception")
    stripe_total = sum(t.get("amount", 0) for t in transactions if t.get("payment_method") in (None, "stripe"))
    mercadopago_total = sum(t.get("amount", 0) for t in transactions if t.get("payment_method") == "mercadopago")
    
    pos_revenue = sum(s.get("total", 0) for s in pos_sales)
    pos_cash = sum(s.get("total", 0) for s in pos_sales if s.get("payment_method") == "cash")
    pos_card = sum(s.get("total", 0) for s in pos_sales if s.get("payment_method") == "card")
    
    total_withdrawals = sum(w.get("amount", 0) for w in withdrawals)
    
    daily_revenue = {}
    for t in transactions:
        date_key = (t.get("created_at") or "")[:10]
        if date_key:
            daily_revenue[date_key] = daily_revenue.get(date_key, 0) + (t.get("amount") or 0)
    for s in pos_sales:
        date_key = (s.get("created_at") or "")[:10]
        if date_key:
            daily_revenue[date_key] = daily_revenue.get(date_key, 0) + (s.get("total") or 0)
    
    daily_chart = [{"date": k, "amount": v} for k, v in sorted(daily_revenue.items())]

    # Revenue breakdown by activity/category
    category_breakdown = {}
    for t in transactions:
        cat = t.get("category") or "General"
        if cat not in category_breakdown:
            category_breakdown[cat] = {"count": 0, "amount": 0.0}
        category_breakdown[cat]["count"] += 1
        category_breakdown[cat]["amount"] += t.get("amount", 0) or 0
    category_rows = [
        {"category": k, "count": v["count"], "amount": v["amount"]}
        for k, v in sorted(category_breakdown.items(), key=lambda x: x[1]["amount"], reverse=True)
    ]
    # Available activities = resolved values from plan_cat_map (plan name or custom category)
    available_categories = sorted(set(plan_cat_map.values())) if plan_cat_map else []

    return {
        "transactions": transactions,
        "pos_sales": pos_sales,
        "withdrawals": withdrawals,
        "summary": {
            "total_revenue": total_revenue + pos_revenue,
            "membership_revenue": total_revenue,
            "pos_revenue": pos_revenue,
            "total_transactions": len(transactions) + len(pos_sales),
            "cash_total": cash_total + pos_cash,
            "card_total": card_total + pos_card,
            "stripe_total": stripe_total,
            "mercadopago_total": mercadopago_total,
            "total_withdrawals": total_withdrawals,
            "net_cash": cash_total + pos_cash - total_withdrawals,
        },
        "daily_chart": daily_chart,
        "category_breakdown": category_rows,
        "available_categories": available_categories,
    }

@router.post("/accounting/withdrawal")
async def create_cash_withdrawal(withdrawal: CashWithdrawalCreate, admin: dict = Depends(get_current_admin)):
    from auth import check_role
    check_role(admin, ["super_admin", "gym_admin", "gym_manager"])
    withdrawal_dict = {
        "id": str(uuid.uuid4()),
        "gym_id": withdrawal.gym_id,
        "amount": withdrawal.amount,
        "reason": withdrawal.reason,
        "notes": withdrawal.notes,
        "registered_by": admin["id"],
        "registered_by_name": admin.get("name"),
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    await db.cash_withdrawals.insert_one(withdrawal_dict)
    withdrawal_dict.pop("_id", None)
    return withdrawal_dict

@router.get("/accounting/withdrawals")
async def get_cash_withdrawals(gym_id: Optional[str] = None, admin: dict = Depends(get_current_admin)):
    query = {}
    if admin["role"] != "super_admin":
        query["gym_id"] = admin.get("gym_id")
    elif gym_id:
        query["gym_id"] = gym_id
    withdrawals = await db.cash_withdrawals.find(query, {"_id": 0}).sort("created_at", -1).to_list(500)
    return withdrawals

@router.get("/accounting/transactions")
async def get_all_transactions(
    gym_id: Optional[str] = None, status: Optional[str] = None,
    date_from: Optional[str] = None, date_to: Optional[str] = None,
    admin: dict = Depends(get_current_admin)
):
    """Get all transactions with paid/pending status for detailed view"""
    query = {}
    if admin["role"] != "super_admin":
        query["gym_id"] = admin.get("gym_id")
    elif gym_id:
        query["gym_id"] = gym_id
    if status:
        query["payment_status"] = status
    if date_from:
        query["created_at"] = {"$gte": date_from}
    if date_to:
        if "created_at" in query:
            query["created_at"]["$lte"] = date_to + "T23:59:59"
        else:
            query["created_at"] = {"$lte": date_to + "T23:59:59"}
    transactions = await db.payment_transactions.find(query, {"_id": 0}).sort("created_at", -1).to_list(1000)
    # Enrich with member data
    for t in transactions:
        if t.get("member_id") and not t.get("member_name"):
            member = await db.members.find_one({"id": t["member_id"]}, {"_id": 0, "name": 1, "code": 1})
            if member:
                t["member_name"] = member.get("name")
                t["member_code"] = member.get("code")
        if t.get("plan_id") and not t.get("plan_name"):
            plan = await db.plans.find_one({"id": t["plan_id"]}, {"_id": 0, "name": 1})
            if plan:
                t["plan_name"] = plan.get("name")
    paid_count = len([t for t in transactions if t.get("payment_status") == "paid"])
    pending_count = len([t for t in transactions if t.get("payment_status") == "pending"])
    total_paid = sum(t.get("amount", 0) for t in transactions if t.get("payment_status") == "paid")
    total_pending = sum(t.get("amount", 0) for t in transactions if t.get("payment_status") == "pending")
    return {
        "transactions": transactions,
        "summary": {
            "paid_count": paid_count,
            "pending_count": pending_count,
            "total_paid": total_paid,
            "total_pending": total_pending,
            "total_count": len(transactions)
        }
    }


@router.delete("/accounting/prune")
async def prune_old_records(months: int = 6, admin: dict = Depends(get_current_admin)):
    from auth import check_role
    check_role(admin, ["super_admin", "gym_admin"])
    from datetime import timedelta
    cutoff = (datetime.now(timezone.utc) - timedelta(days=months * 30)).isoformat()
    gym_id = admin.get("gym_id")
    query = {"created_at": {"$lt": cutoff}}
    if gym_id:
        query["gym_id"] = gym_id
    
    access_deleted = await db.access_logs.delete_many(query)
    
    log_query = {"timestamp": {"$lt": cutoff}}
    if gym_id:
        log_query["gym_id"] = gym_id
    access_log_deleted = await db.access_logs.delete_many(log_query)
    
    return {
        "message": f"Registros anteriores a {months} meses eliminados",
        "access_logs_deleted": access_deleted.deleted_count + access_log_deleted.deleted_count,
    }

@router.get("/accounting/pdf")
async def generate_accounting_pdf(
    gym_id: Optional[str] = None, date_from: Optional[str] = None,
    date_to: Optional[str] = None, admin: dict = Depends(get_current_admin)
):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from io import BytesIO
    
    query = {"payment_status": "paid"}
    if admin["role"] != "super_admin":
        query["gym_id"] = admin.get("gym_id")
    elif gym_id:
        query["gym_id"] = gym_id
    if date_from:
        query["created_at"] = {"$gte": date_from}
    if date_to:
        if "created_at" in query:
            query["created_at"]["$lte"] = date_to + "T23:59:59"
        else:
            query["created_at"] = {"$lte": date_to + "T23:59:59"}
    transactions = await db.payment_transactions.find(query, {"_id": 0}).sort("created_at", 1).to_list(5000)
    gym_name = "Todos los Gimnasios"
    target_gym_id = admin.get("gym_id") or gym_id
    if target_gym_id:
        gym_doc = await db.gyms.find_one({"id": target_gym_id}, {"_id": 0, "name": 1})
        gym_name = gym_doc.get("name", gym_name) if gym_doc else gym_name
    total = sum(t.get("amount", 0) for t in transactions)
    
    # Get membership breakdown (users per plan)
    target_gym_id2 = admin.get("gym_id") or gym_id
    ms_query = {"status": "active"}
    if target_gym_id2:
        ms_query["gym_id"] = target_gym_id2
    active_memberships = await db.memberships.find(ms_query, {"_id": 0, "plan_id": 1}).to_list(10000)
    plan_counts = {}
    for ms in active_memberships:
        pid = ms.get("plan_id", "")
        plan_counts[pid] = plan_counts.get(pid, 0) + 1
    plan_names = {}
    for pid in plan_counts:
        p = await db.plans.find_one({"id": pid}, {"_id": 0, "name": 1, "price": 1})
        if p:
            plan_names[pid] = {"name": p.get("name", "Sin nombre"), "price": p.get("price", 0)}
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=20*mm, bottomMargin=20*mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('CustomTitle', parent=styles['Title'], fontSize=18, spaceAfter=6)
    subtitle_style = ParagraphStyle('CustomSubtitle', parent=styles['Normal'], fontSize=10, textColor=colors.grey)
    elements = []
    elements.append(Paragraph(f"Informe de Contabilidad - {gym_name}", title_style))
    period = ""
    if date_from:
        period += f"Desde: {date_from} "
    if date_to:
        period += f"Hasta: {date_to}"
    if not period:
        period = "Todos los periodos"
    elements.append(Paragraph(period, subtitle_style))
    elements.append(Spacer(1, 10*mm))
    elements.append(Paragraph(f"<b>Total Recaudado: ${total:,.2f}</b>  |  Transacciones: {len(transactions)}", styles['Normal']))
    elements.append(Spacer(1, 8*mm))
    
    # Membership breakdown table
    if plan_counts:
        section_style = ParagraphStyle('Sec', parent=styles['Heading2'], fontSize=12, spaceBefore=4, spaceAfter=4, textColor=colors.HexColor('#18181B'))
        elements.append(Paragraph("Usuarios por Membresia (activos)", section_style))
        mb_data = [['Plan', 'Precio', 'Usuarios Activos', 'Ingreso Estimado']]
        total_active = 0
        total_estimated = 0
        for pid, count in sorted(plan_counts.items(), key=lambda x: x[1], reverse=True):
            info = plan_names.get(pid, {"name": "Desconocido", "price": 0})
            estimated = info["price"] * count
            total_active += count
            total_estimated += estimated
            mb_data.append([info["name"], f'${info["price"]:,.2f}', str(count), f'${estimated:,.2f}'])
        mb_data.append(['TOTAL', '', str(total_active), f'${total_estimated:,.2f}'])
        mb_table = Table(mb_data, colWidths=[150, 70, 100, 100])
        mb_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A5F')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.lightgrey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F0F7FF')]),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E3F2FD')),
            ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
            ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(mb_table)
        elements.append(Spacer(1, 8*mm))
    
    if transactions:
        table_data = [['Fecha', 'Socio', 'Plan', 'Metodo', 'Monto']]
        for t in transactions:
            method_map = {"cash": "Efectivo", "card_reception": "Tarjeta", "stripe": "Stripe Online", "mercadopago": "MercadoPago", "redsys": "Redsys TPV"}
            method = method_map.get(t.get("payment_method", "stripe"), "Stripe Online")
            date_str = t.get("created_at", "")[:10]
            table_data.append([date_str, t.get("member_name", "-"), t.get("plan_name", "-"), method, f"${t.get('amount', 0):,.2f}"])
        col_widths = [70, 130, 100, 80, 70]
        table = Table(table_data, colWidths=col_widths)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#18181B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, 0), 9), ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('ALIGN', (-1, 0), (-1, -1), 'RIGHT'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.lightgrey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F5F5')]),
            ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(table)
    else:
        elements.append(Paragraph("No hay transacciones en este periodo.", styles['Normal']))
    doc.build(elements)
    buffer.seek(0)
    filename = f"contabilidad_{gym_name.replace(' ', '_')}_{date_from or 'all'}_{date_to or 'all'}.pdf"
    return StreamingResponse(buffer, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename={filename}"})


@router.get("/accounting/excel")
async def generate_accounting_excel(
    gym_id: Optional[str] = None, date_from: Optional[str] = None,
    date_to: Optional[str] = None, status: Optional[str] = None,
    category: Optional[str] = None,
    admin: dict = Depends(get_current_admin)
):
    """Export detailed Excel with memberships, payments, status, and summary."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from io import BytesIO

    target_gym_id = admin.get("gym_id") or gym_id

    # Fetch transactions (all statuses)
    tx_query = {}
    if target_gym_id:
        tx_query["gym_id"] = target_gym_id
    if status:
        tx_query["payment_status"] = status
    if date_from:
        tx_query["created_at"] = {"$gte": date_from}
    if date_to:
        if "created_at" in tx_query:
            tx_query["created_at"]["$lte"] = date_to + "T23:59:59"
        else:
            tx_query["created_at"] = {"$lte": date_to + "T23:59:59"}

    transactions = await db.payment_transactions.find(tx_query, {"_id": 0}).sort("created_at", -1).to_list(10000)

    # Load plan -> activity map (plan's custom category if set, otherwise plan name)
    plan_q = {}
    if target_gym_id:
        plan_q["gym_id"] = target_gym_id
    plans_all = await db.plans.find(plan_q, {"_id": 0, "id": 1, "name": 1, "category": 1, "price": 1, "duration_days": 1}).to_list(500)
    plan_cat_map = {}
    for p in plans_all:
        cat = p.get("category")
        if not cat or cat.strip().lower() == "general":
            cat = p.get("name") or "Sin plan"
        plan_cat_map[p["id"]] = cat

    # Attach activity to each tx
    for t in transactions:
        activity = plan_cat_map.get(t.get("plan_id"))
        if not activity:
            activity = t.get("plan_name") or "Sin plan"
        t["category"] = activity

    # Optional activity filter
    if category and category != "all":
        transactions = [t for t in transactions if (t.get("category") or "Sin plan") == category]

    # Enrich with member + plan data
    for t in transactions:
        if t.get("member_id") and not t.get("member_name"):
            member = await db.members.find_one({"id": t["member_id"]}, {"_id": 0, "name": 1, "code": 1, "email": 1, "phone": 1})
            if member:
                t["member_name"] = member.get("name")
                t["member_code"] = member.get("code")
                t["member_email"] = member.get("email", "")
                t["member_phone"] = member.get("phone", "")
        if t.get("plan_id") and not t.get("plan_name"):
            plan = await db.plans.find_one({"id": t["plan_id"]}, {"_id": 0, "name": 1, "price": 1, "duration_days": 1})
            if plan:
                t["plan_name"] = plan.get("name")
                t["plan_price"] = plan.get("price")
                t["plan_duration"] = plan.get("duration_days")

    # Fetch memberships for member status
    member_ids = list(set(t.get("member_id") for t in transactions if t.get("member_id")))
    memberships = {}
    for mid in member_ids:
        ms = await db.memberships.find_one(
            {"member_id": mid, "status": {"$in": ["active", "expired"]}},
            {"_id": 0, "status": 1, "end_date": 1, "start_date": 1}
        )
        if ms:
            memberships[mid] = ms

    # POS sales
    pos_query = {}
    if target_gym_id:
        pos_query["gym_id"] = target_gym_id
    if date_from:
        pos_query["created_at"] = {"$gte": date_from}
    if date_to:
        if "created_at" in pos_query:
            pos_query["created_at"]["$lte"] = date_to + "T23:59:59"
        else:
            pos_query["created_at"] = {"$lte": date_to + "T23:59:59"}
    pos_sales = await db.pos_sales.find(pos_query, {"_id": 0}).sort("created_at", -1).to_list(5000)

    # Withdrawals
    w_query = {}
    if target_gym_id:
        w_query["gym_id"] = target_gym_id
    if date_from:
        w_query["created_at"] = {"$gte": date_from}
    if date_to:
        if "created_at" in w_query:
            w_query["created_at"]["$lte"] = date_to + "T23:59:59"
        else:
            w_query["created_at"] = {"$lte": date_to + "T23:59:59"}
    withdrawals = await db.cash_withdrawals.find(w_query, {"_id": 0}).sort("created_at", -1).to_list(1000)

    # Gym name
    gym_name = "Todos"
    if target_gym_id:
        gym_doc = await db.gyms.find_one({"id": target_gym_id}, {"_id": 0, "name": 1})
        gym_name = gym_doc.get("name", "Gym") if gym_doc else "Gym"

    # Create workbook
    wb = Workbook()

    # --- Sheet 1: RESUMEN ---
    ws_summary = wb.active
    ws_summary.title = "Resumen"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="18181B", end_color="18181B", fill_type="solid")
    green_fill = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid")
    red_fill = PatternFill(start_color="FADBD8", end_color="FADBD8", fill_type="solid")
    amber_fill = PatternFill(start_color="FEF9E7", end_color="FEF9E7", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='D5D5D5'), right=Side(style='thin', color='D5D5D5'),
        top=Side(style='thin', color='D5D5D5'), bottom=Side(style='thin', color='D5D5D5')
    )

    ws_summary.append([f"INFORME DE CONTABILIDAD - {gym_name.upper()}"])
    ws_summary.merge_cells('A1:F1')
    ws_summary['A1'].font = Font(bold=True, size=14)
    ws_summary.append([f"Periodo: {date_from or 'Inicio'} a {date_to or 'Hoy'}"])
    ws_summary.append([])

    paid = [t for t in transactions if t.get("payment_status") == "paid"]
    pending = [t for t in transactions if t.get("payment_status") != "paid"]
    total_paid = sum(t.get("amount", 0) for t in paid)
    total_pending = sum(t.get("amount", 0) for t in pending)
    total_pos = sum(s.get("total", 0) for s in pos_sales)
    total_withdrawals_amount = sum(w.get("amount", 0) for w in withdrawals)

    summary_data = [
        ["Concepto", "Cantidad", "Monto"],
        ["Pagos Completados", len(paid), total_paid],
        ["Pagos Pendientes", len(pending), total_pending],
        ["Total Membresias", len(transactions), total_paid + total_pending],
        ["Ventas POS", len(pos_sales), total_pos],
        ["TOTAL INGRESOS", len(transactions) + len(pos_sales), total_paid + total_pos],
        ["Retiros de Caja", len(withdrawals), -total_withdrawals_amount],
        ["NETO", "", total_paid + total_pos - total_withdrawals_amount],
    ]
    for row in summary_data:
        ws_summary.append(row)

    for col in range(1, 4):
        ws_summary.cell(row=4, column=col).font = header_font
        ws_summary.cell(row=4, column=col).fill = header_fill
    for row_idx in range(5, 12):
        for col in range(1, 4):
            ws_summary.cell(row=row_idx, column=col).border = thin_border
    ws_summary.cell(row=5, column=3).number_format = '#,##0.00'
    ws_summary.cell(row=6, column=3).number_format = '#,##0.00'
    ws_summary.cell(row=6, column=1).fill = amber_fill
    ws_summary.cell(row=6, column=2).fill = amber_fill
    ws_summary.cell(row=6, column=3).fill = amber_fill
    for col in range(1, 4):
        ws_summary.cell(row=9, column=col).font = Font(bold=True)
        ws_summary.cell(row=9, column=col).fill = green_fill
        ws_summary.cell(row=11, column=col).font = Font(bold=True, size=12)

    # By payment method
    ws_summary.append([])
    ws_summary.append(["DESGLOSE POR METODO DE PAGO", "", ""])
    method_map = {"cash": "Efectivo", "card_reception": "Tarjeta Recepcion", "stripe": "Stripe Online", "mercadopago": "MercadoPago", "redsys": "Redsys"}
    methods = {}
    for t in paid:
        m = method_map.get(t.get("payment_method", "stripe"), "Stripe Online")
        methods[m] = methods.get(m, {"count": 0, "amount": 0})
        methods[m]["count"] += 1
        methods[m]["amount"] += t.get("amount", 0)
    ws_summary.append(["Metodo", "Transacciones", "Monto"])
    for m, v in methods.items():
        ws_summary.append([m, v["count"], v["amount"]])

    ws_summary.column_dimensions['A'].width = 30
    ws_summary.column_dimensions['B'].width = 15
    ws_summary.column_dimensions['C'].width = 18

    # Membership breakdown: users per plan
    ms_query2 = {"status": "active"}
    if target_gym_id:
        ms_query2["gym_id"] = target_gym_id
    active_ms = await db.memberships.find(ms_query2, {"_id": 0, "plan_id": 1}).to_list(10000)
    plan_counts = {}
    for ms in active_ms:
        pid = ms.get("plan_id", "")
        plan_counts[pid] = plan_counts.get(pid, 0) + 1
    
    if plan_counts:
        ws_summary.append([])
        ws_summary.append([])
        r = ws_summary.max_row + 1
        ws_summary.append(["USUARIOS POR MEMBRESIA (ACTIVOS)", "", "", ""])
        ws_summary.cell(row=ws_summary.max_row, column=1).font = Font(bold=True, size=12)
        ws_summary.append(["Plan", "Precio", "Usuarios Activos", "Ingreso Estimado"])
        for col in range(1, 5):
            ws_summary.cell(row=ws_summary.max_row, column=col).font = header_font
            ws_summary.cell(row=ws_summary.max_row, column=col).fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
        
        total_active = 0
        total_estimated = 0
        for pid, count in sorted(plan_counts.items(), key=lambda x: x[1], reverse=True):
            plan_doc = await db.plans.find_one({"id": pid}, {"_id": 0, "name": 1, "price": 1})
            pname = plan_doc.get("name", "Desconocido") if plan_doc else "Desconocido"
            pprice = plan_doc.get("price", 0) if plan_doc else 0
            estimated = pprice * count
            total_active += count
            total_estimated += estimated
            ws_summary.append([pname, pprice, count, estimated])
            ws_summary.cell(row=ws_summary.max_row, column=2).number_format = '#,##0.00'
            ws_summary.cell(row=ws_summary.max_row, column=4).number_format = '#,##0.00'
        
        ws_summary.append(["TOTAL", "", total_active, total_estimated])
        ws_summary.cell(row=ws_summary.max_row, column=1).font = Font(bold=True)
        ws_summary.cell(row=ws_summary.max_row, column=3).font = Font(bold=True)
        ws_summary.cell(row=ws_summary.max_row, column=4).font = Font(bold=True)
        ws_summary.cell(row=ws_summary.max_row, column=4).number_format = '#,##0.00'
        for col in range(1, 5):
            ws_summary.cell(row=ws_summary.max_row, column=col).fill = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
    
    ws_summary.column_dimensions['D'].width = 20

    # --- Sheet 2: PAGOS DETALLADOS ---
    ws_tx = wb.create_sheet("Pagos Membresias")
    headers = ["Fecha", "Hora", "Socio", "Codigo", "Email", "Telefono", "Plan", "Actividad", "Duracion (dias)", "Metodo Pago", "Estado", "Monto", "Estado Membresia", "Vencimiento"]
    ws_tx.append(headers)
    for col_idx, h in enumerate(headers, 1):
        cell = ws_tx.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for t in transactions:
        ms = memberships.get(t.get("member_id"), {})
        status_label = "Pagado" if t.get("payment_status") == "paid" else "Pendiente"
        ms_status = (ms.get("status") or "").capitalize() if ms else ""
        ms_end = (ms.get("end_date") or "")[:10] if ms else ""
        method = method_map.get(t.get("payment_method", "stripe"), "Stripe Online")
        created = t.get("created_at") or ""
        row = [
            created[:10],
            created[11:19] if len(created) > 19 else "",
            t.get("member_name") or "",
            t.get("member_code") or "",
            t.get("member_email") or "",
            t.get("member_phone") or "",
            t.get("plan_name") or "",
            t.get("category") or "General",
            t.get("plan_duration") or "",
            method,
            status_label,
            t.get("amount") or 0,
            ms_status,
            ms_end,
        ]
        ws_tx.append(row)

    # Color rows by status
    for row_idx in range(2, len(transactions) + 2):
        status_cell = ws_tx.cell(row=row_idx, column=11)
        if status_cell.value == "Pagado":
            status_cell.fill = green_fill
        else:
            status_cell.fill = amber_fill
        ws_tx.cell(row=row_idx, column=12).number_format = '#,##0.00'
        for col in range(1, 15):
            ws_tx.cell(row=row_idx, column=col).border = thin_border

    widths = [12, 10, 25, 10, 25, 15, 20, 16, 14, 18, 12, 12, 16, 14]
    for i, w in enumerate(widths, 1):
        col_letter = chr(64 + i) if i <= 26 else 'A'
        ws_tx.column_dimensions[col_letter].width = w

    # --- Sheet 2b: DESGLOSE POR ACTIVIDAD ---
    cat_totals = {}
    for t in transactions:
        if t.get("payment_status") != "paid":
            continue
        c = t.get("category") or "General"
        if c not in cat_totals:
            cat_totals[c] = {"count": 0, "amount": 0.0}
        cat_totals[c]["count"] += 1
        cat_totals[c]["amount"] += t.get("amount", 0) or 0
    if cat_totals:
        ws_cat = wb.create_sheet("Por Actividad")
        ws_cat.append(["Actividad", "Transacciones", "Ingresos"])
        for col_idx in range(1, 4):
            cell = ws_cat.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
        for cat_name, v in sorted(cat_totals.items(), key=lambda x: x[1]["amount"], reverse=True):
            ws_cat.append([cat_name, v["count"], v["amount"]])
            ws_cat.cell(row=ws_cat.max_row, column=3).number_format = '#,##0.00'
        total_count = sum(v["count"] for v in cat_totals.values())
        total_amount = sum(v["amount"] for v in cat_totals.values())
        ws_cat.append(["TOTAL", total_count, total_amount])
        for col_idx in range(1, 4):
            ws_cat.cell(row=ws_cat.max_row, column=col_idx).font = Font(bold=True)
            ws_cat.cell(row=ws_cat.max_row, column=col_idx).fill = green_fill
        ws_cat.cell(row=ws_cat.max_row, column=3).number_format = '#,##0.00'
        ws_cat.column_dimensions['A'].width = 25
        ws_cat.column_dimensions['B'].width = 16
        ws_cat.column_dimensions['C'].width = 18

    # --- Sheet 3: VENTAS POS ---
    if pos_sales:
        ws_pos = wb.create_sheet("Ventas POS")
        pos_headers = ["Fecha", "Hora", "Productos", "Metodo", "Total"]
        ws_pos.append(pos_headers)
        for col_idx, h in enumerate(pos_headers, 1):
            cell = ws_pos.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
        for s in pos_sales:
            items_str = ", ".join([f"{i.get('product_name','?')} x{i.get('quantity',1)}" for i in (s.get("items") or [])])
            s_created = s.get("created_at") or ""
            ws_pos.append([
                s_created[:10],
                s_created[11:19] if len(s_created) > 19 else "",
                items_str,
                "Efectivo" if s.get("payment_method") == "cash" else "Tarjeta",
                s.get("total") or 0,
            ])
        ws_pos.column_dimensions['A'].width = 12
        ws_pos.column_dimensions['B'].width = 10
        ws_pos.column_dimensions['C'].width = 50
        ws_pos.column_dimensions['D'].width = 14
        ws_pos.column_dimensions['E'].width = 12

    # --- Sheet 4: RETIROS ---
    if withdrawals:
        ws_w = wb.create_sheet("Retiros de Caja")
        w_headers = ["Fecha", "Hora", "Motivo", "Notas", "Responsable", "Monto"]
        ws_w.append(w_headers)
        for col_idx, h in enumerate(w_headers, 1):
            cell = ws_w.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
        for w in withdrawals:
            w_created = w.get("created_at") or ""
            ws_w.append([
                w_created[:10],
                w_created[11:19] if len(w_created) > 19 else "",
                w.get("reason") or "",
                w.get("notes") or "",
                w.get("registered_by_name") or "",
                w.get("amount") or 0,
            ])
            ws_w.cell(row=ws_w.max_row, column=6).fill = red_fill
        ws_w.column_dimensions['A'].width = 12
        ws_w.column_dimensions['B'].width = 10
        ws_w.column_dimensions['C'].width = 30
        ws_w.column_dimensions['D'].width = 30
        ws_w.column_dimensions['E'].width = 20
        ws_w.column_dimensions['F'].width = 12

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    filename = f"contabilidad_{gym_name.replace(' ', '_')}_{date_from or 'all'}_{date_to or 'all'}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
@router.get("/accounting/sales-report-pdf")
async def generate_sales_report_pdf(
    gym_id: Optional[str] = None, period: str = "daily",
    date_from: Optional[str] = None, date_to: Optional[str] = None,
    admin: dict = Depends(get_current_admin)
):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from io import BytesIO
    from datetime import timedelta
    
    now = datetime.now(timezone.utc)
    if not date_from:
        if period == "daily":
            date_from = now.replace(hour=0, minute=0, second=0, microsecond=0).isoformat()[:10]
        elif period == "weekly":
            date_from = (now - timedelta(days=now.weekday())).replace(hour=0, minute=0, second=0, microsecond=0).isoformat()[:10]
        elif period == "monthly":
            date_from = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0).isoformat()[:10]
        else:
            date_from = (now - timedelta(days=365)).isoformat()[:10]
    if not date_to:
        date_to = now.isoformat()[:10]
    
    target_gym_id = admin.get("gym_id") or gym_id
    
    # Membership payments
    pay_query = {"payment_status": "paid", "created_at": {"$gte": date_from, "$lte": date_to + "T23:59:59"}}
    if target_gym_id:
        pay_query["gym_id"] = target_gym_id
    transactions = await db.payment_transactions.find(pay_query, {"_id": 0}).sort("created_at", 1).to_list(5000)
    
    # POS sales
    pos_query = {"created_at": {"$gte": date_from, "$lte": date_to + "T23:59:59"}}
    if target_gym_id:
        pos_query["gym_id"] = target_gym_id
    pos_sales = await db.pos_sales.find(pos_query, {"_id": 0}).sort("created_at", 1).to_list(5000)
    
    # Withdrawals
    w_query = {"created_at": {"$gte": date_from, "$lte": date_to + "T23:59:59"}}
    if target_gym_id:
        w_query["gym_id"] = target_gym_id
    withdrawals = await db.cash_withdrawals.find(w_query, {"_id": 0}).sort("created_at", 1).to_list(1000)
    
    gym_name = "Todos los Gimnasios"
    if target_gym_id:
        gym_doc = await db.gyms.find_one({"id": target_gym_id}, {"_id": 0, "name": 1})
        gym_name = gym_doc.get("name", gym_name) if gym_doc else gym_name
    
    total_memberships = sum(t.get("amount", 0) for t in transactions)
    total_pos = sum(s.get("total", 0) for s in pos_sales)
    total_withdrawals = sum(w.get("amount", 0) for w in withdrawals)
    grand_total = total_memberships + total_pos
    
    period_labels = {"daily": "Diario", "weekly": "Semanal", "monthly": "Mensual", "custom": "Personalizado"}
    
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=15*mm, bottomMargin=15*mm, leftMargin=12*mm, rightMargin=12*mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('T', parent=styles['Title'], fontSize=16, spaceAfter=4, textColor=colors.HexColor('#18181B'))
    sub_style = ParagraphStyle('S', parent=styles['Normal'], fontSize=9, textColor=colors.grey)
    section_style = ParagraphStyle('Sec', parent=styles['Heading2'], fontSize=12, spaceBefore=10, spaceAfter=4, textColor=colors.HexColor('#18181B'))
    
    elements = []
    elements.append(Paragraph(f"Informe de Ventas - {gym_name}", title_style))
    elements.append(Paragraph(f"Periodo: {period_labels.get(period, period)} | Desde: {date_from} | Hasta: {date_to}", sub_style))
    elements.append(Spacer(1, 6*mm))
    
    # Summary table
    summary_data = [
        ['Concepto', 'Total'],
        ['Ingresos por Membresias', f'${total_memberships:,.2f}'],
        ['Ingresos por Ventas POS', f'${total_pos:,.2f}'],
        ['TOTAL INGRESOS', f'${grand_total:,.2f}'],
        ['Retiros de Caja', f'-${total_withdrawals:,.2f}'],
        ['NETO', f'${(grand_total - total_withdrawals):,.2f}'],
    ]
    t = Table(summary_data, colWidths=[300, 150])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#18181B')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ('BACKGROUND', (0, 3), (-1, 3), colors.HexColor('#E8F5E9')),
        ('BACKGROUND', (0, 5), (-1, 5), colors.HexColor('#E3F2FD')),
        ('FONTNAME', (0, 3), (-1, 3), 'Helvetica-Bold'),
        ('FONTNAME', (0, 5), (-1, 5), 'Helvetica-Bold'),
        ('TOPPADDING', (0, 0), (-1, -1), 4), ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(t)
    elements.append(Spacer(1, 8*mm))
    
    # Memberships detail
    if transactions:
        elements.append(Paragraph("Detalle de Membresias", section_style))
        mem_data = [['Fecha', 'Socio', 'Plan', 'Metodo', 'Monto']]
        method_map = {"cash": "Efectivo", "card_reception": "Tarjeta", "stripe": "Stripe", "mercadopago": "MercadoPago", "redsys": "Redsys TPV"}
        for tx in transactions:
            mem_data.append([
                tx.get("created_at", "")[:10],
                tx.get("member_name", "-"),
                tx.get("plan_name", "-"),
                method_map.get(tx.get("payment_method", "stripe"), "Stripe"),
                f'${tx.get("amount", 0):,.2f}'
            ])
        mt = Table(mem_data, colWidths=[65, 130, 110, 75, 70])
        mt.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#18181B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (-1, 0), (-1, -1), 'RIGHT'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.lightgrey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FAFAFA')]),
            ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        elements.append(mt)
        elements.append(Spacer(1, 6*mm))
    
    # POS detail
    if pos_sales:
        elements.append(Paragraph("Detalle de Ventas POS", section_style))
        pos_data = [['Fecha', 'Productos', 'Metodo', 'Total']]
        for s in pos_sales:
            items_str = ", ".join([f"{i['product_name']}({i['quantity']})" for i in s.get("items", [])])
            pos_data.append([
                s.get("created_at", "")[:10],
                items_str[:50],
                "Efectivo" if s.get("payment_method") == "cash" else "Tarjeta",
                f'${s.get("total", 0):,.2f}'
            ])
        pt = Table(pos_data, colWidths=[65, 220, 75, 70])
        pt.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#18181B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (-1, 0), (-1, -1), 'RIGHT'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.lightgrey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FAFAFA')]),
            ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        elements.append(pt)
        elements.append(Spacer(1, 6*mm))
    
    # Withdrawals detail
    if withdrawals:
        elements.append(Paragraph("Detalle de Retiros de Caja", section_style))
        wd = [['Fecha', 'Motivo', 'Responsable', 'Monto']]
        for w in withdrawals:
            wd.append([w.get("created_at", "")[:10], w.get("reason", ""), w.get("registered_by_name", ""), f'-${w.get("amount", 0):,.2f}'])
        wt = Table(wd, colWidths=[65, 190, 110, 70])
        wt.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#18181B')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('ALIGN', (-1, 0), (-1, -1), 'RIGHT'),
            ('TEXTCOLOR', (-1, 1), (-1, -1), colors.red),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.lightgrey),
            ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        elements.append(wt)
    
    doc.build(elements)
    buffer.seek(0)
    filename = f"informe_ventas_{period}_{gym_name.replace(' ','_')}_{date_from}_{date_to}.pdf"
    return StreamingResponse(buffer, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={filename}"})
