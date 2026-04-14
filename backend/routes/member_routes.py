from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from datetime import datetime, timezone, timedelta
from typing import Optional
import uuid
import logging
import io

from database import db
from auth import get_current_admin, create_jwt_token, check_role, check_permission
from models import MemberCreate, MemberPublicRegister, MemberUpdate
from qr_utils import generate_member_code

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api")

@router.post("/members")
async def create_member(member: MemberCreate, admin: dict = Depends(get_current_admin)):
    check_permission(admin, "members_create")
    existing = await db.members.find_one({"email": member.email, "gym_id": member.gym_id})
    if existing:
        raise HTTPException(status_code=400, detail="Email already registered in this gym")
    member_dict = member.model_dump()
    member_dict["id"] = str(uuid.uuid4())
    member_dict["code"] = generate_member_code()
    member_dict["status"] = "active"
    member_dict["gender"] = member.gender or "prefer_not_to_say"
    member_dict["created_at"] = datetime.now(timezone.utc).isoformat()
    while await db.members.find_one({"code": member_dict["code"]}):
        member_dict["code"] = generate_member_code()
    await db.members.insert_one(member_dict)
    member_dict.pop("_id", None)
    return member_dict

@router.post("/members/register")
async def register_member_public(member: MemberPublicRegister):
    existing = await db.members.find_one({"email": member.email, "gym_id": member.gym_id})
    if existing:
        raise HTTPException(status_code=400, detail="Este email ya esta registrado en este gimnasio")
    gym = await db.gyms.find_one({"id": member.gym_id})
    if not gym:
        raise HTTPException(status_code=404, detail="Gimnasio no encontrado")
    if gym.get("status") == "suspended":
        raise HTTPException(status_code=403, detail="Este gimnasio esta suspendido")
    
    # If a plan is selected, member starts as "pending" until payment
    initial_status = "pending" if member.plan_id else "active"
    
    member_dict = {
        "id": str(uuid.uuid4()),
        "email": member.email,
        "name": member.name,
        "phone": member.phone,
        "gym_id": member.gym_id,
        "code": generate_member_code(),
        "status": initial_status,
        "gender": member.gender or "prefer_not_to_say",
        "form_responses": member.form_responses,
        "created_at": datetime.now(timezone.utc).isoformat()
    }
    while await db.members.find_one({"code": member_dict["code"]}):
        member_dict["code"] = generate_member_code()
    await db.members.insert_one(member_dict)
    member_dict.pop("_id", None)
    
    membership_data = None
    plan_data = None
    if member.plan_id:
        plan = await db.plans.find_one({"id": member.plan_id}, {"_id": 0})
        if plan:
            plan_data = plan
            # Create membership as pending_payment (NOT active)
            membership = {
                "id": str(uuid.uuid4()),
                "member_id": member_dict["id"],
                "plan_id": plan["id"],
                "gym_id": member.gym_id,
                "start_date": None,
                "end_date": None,
                "status": "pending_payment",
                "created_at": datetime.now(timezone.utc).isoformat()
            }
            await db.memberships.insert_one(membership)
            membership.pop("_id", None)
            membership_data = membership
            
            # Create a pending payment transaction for tracking
            transaction = {
                "id": str(uuid.uuid4()),
                "member_id": member_dict["id"],
                "member_name": member_dict["name"],
                "plan_id": plan["id"],
                "plan_name": plan.get("name"),
                "gym_id": member.gym_id,
                "amount": plan.get("price", 0),
                "currency": gym.get("currency", "eur"),
                "payment_method": "pending",
                "status": "pending",
                "payment_status": "pending",
                "created_at": datetime.now(timezone.utc).isoformat()
            }
            await db.payment_transactions.insert_one(transaction)
    
    token = create_jwt_token({"sub": member_dict["id"], "role": "member", "gym_id": member.gym_id})
    
    # Send welcome email with payment button if plan selected
    if member.plan_id and plan_data:
        try:
            from routes.misc_routes import send_gym_email
            gym_color = gym.get("primary_color", "#E1FF01")
            gym_name = gym.get("name", "IngresoQR")
            # Build payment URL
            app_url = "https://app.ingresoqr.com"
            payment_section = f"""
            <div style="background:#18181B;padding:20px;border-radius:12px;margin:20px 0;">
                <p style="color:#a1a1aa;margin:0 0 5px;font-size:14px;">Plan seleccionado</p>
                <p style="font-size:22px;font-weight:700;color:{gym_color};margin:0;">{plan_data.get('name','')}</p>
                <p style="color:#fff;font-size:18px;font-weight:700;margin:8px 0 0;">{plan_data.get('price',0)} {gym.get('currency','EUR').upper()}</p>
                <p style="color:#a1a1aa;margin:5px 0 0;">{plan_data.get('duration_days',30)} dias de acceso</p>
            </div>
            <div style="background:#F59E0B22;border:1px solid #F59E0B55;padding:15px;border-radius:12px;margin:15px 0;">
                <p style="color:#FBBF24;font-weight:700;margin:0 0 5px;font-size:16px;">Pago pendiente</p>
                <p style="color:#E5E5E5;margin:0;font-size:14px;">Para habilitar tu acceso al gimnasio, realiza el pago de tu membresia.</p>
            </div>
            <div style="text-align:center;margin:25px 0;">
                <a href="{app_url}/app/membership" style="display:inline-block;padding:16px 40px;background:#F59E0B;color:#000;font-weight:800;text-decoration:none;border-radius:12px;font-size:18px;">Pagar Ahora</a>
            </div>
            """
            html = f"""
            <div style="font-family:Arial,sans-serif;max-width:500px;margin:0 auto;padding:30px;background:#09090B;color:#fff;border-radius:16px;">
                <div style="text-align:center;margin-bottom:20px;">
                    <h1 style="color:{gym_color};margin:0;">{gym_name}</h1>
                </div>
                <h2 style="margin:0 0 10px;">Bienvenido/a, {member_dict.get('name','Socio')}</h2>
                <p style="color:#a1a1aa;">Tu registro en <strong style="color:#fff;">{gym_name}</strong> ha sido completado.</p>
                <div style="background:#18181B;padding:20px;border-radius:12px;text-align:center;margin:20px 0;">
                    <p style="color:#a1a1aa;margin:0 0 8px;">Tu codigo de acceso</p>
                    <p style="font-size:32px;font-weight:900;color:{gym_color};font-family:monospace;letter-spacing:4px;margin:0;">{member_dict.get('code','------')}</p>
                </div>
                {payment_section}
                <p style="color:#71717A;font-size:12px;text-align:center;margin-top:30px;">Guarda este codigo para acceder al gimnasio una vez realizado el pago.</p>
            </div>
            """
            await send_gym_email(
                member.gym_id, member_dict["email"],
                f"Bienvenido/a a {gym_name} - Completa tu pago",
                html, member_id=member_dict["id"], email_type="welcome_payment"
            )
        except Exception as e:
            logger.warning(f"Could not send welcome payment email: {e}")
    
    return {
        "member": member_dict,
        "membership": membership_data,
        "plan": plan_data,
        "token": token,
        "message": f"Registro exitoso. Tu codigo de acceso es: {member_dict['code']}",
        "requires_payment": bool(member.plan_id)
    }

@router.get("/members")
async def get_members(gym_id: Optional[str] = None, status: Optional[str] = None, admin: dict = Depends(get_current_admin)):
    check_permission(admin, "members_view")
    query = {}
    if admin["role"] != "super_admin":
        query["gym_id"] = admin.get("gym_id")
    elif gym_id:
        query["gym_id"] = gym_id
    if status:
        query["status"] = status
    members = await db.members.find(query, {"_id": 0}).to_list(1000)

    # Enrich with active membership data
    member_ids = [m["id"] for m in members]
    if member_ids:
        memberships = await db.memberships.find(
            {"member_id": {"$in": member_ids}, "status": {"$in": ["active", "expired"]}},
            {"_id": 0}
        ).sort("created_at", -1).to_list(5000)

        plan_ids = list(set(m.get("plan_id") for m in memberships if m.get("plan_id")))
        plans_map = {}
        if plan_ids:
            plans_list = await db.plans.find({"id": {"$in": plan_ids}}, {"_id": 0}).to_list(500)
            plans_map = {p["id"]: p for p in plans_list}

        membership_map = {}
        for ms in memberships:
            mid = ms["member_id"]
            if mid not in membership_map:
                membership_map[mid] = ms

        for member in members:
            ms = membership_map.get(member["id"])
            if ms:
                plan = plans_map.get(ms.get("plan_id"))
                member["membership"] = {
                    "plan_name": plan.get("name") if plan else None,
                    "start_date": ms.get("start_date"),
                    "end_date": ms.get("end_date"),
                    "status": ms.get("status"),
                    "payment_status": ms.get("payment_status"),
                }
            else:
                member["membership"] = None

    return members

@router.put("/members/{member_id}/membership")
async def update_member_membership(member_id: str, data: dict, admin: dict = Depends(get_current_admin)):
    """Update membership expiration date with comment. Gym admin or manager with membership_edit permission."""
    check_permission(admin, "membership_edit")

    member = await db.members.find_one({"id": member_id}, {"_id": 0})
    if not member:
        raise HTTPException(status_code=404, detail="Socio no encontrado")

    if admin["role"] != "super_admin" and member.get("gym_id") != admin.get("gym_id"):
        raise HTTPException(status_code=403, detail="No tienes acceso a este socio")

    new_end_date = data.get("end_date")
    comment = data.get("comment", "").strip()

    if not new_end_date:
        raise HTTPException(status_code=400, detail="Fecha de vencimiento requerida")

    # Find active membership
    membership = await db.memberships.find_one(
        {"member_id": member_id, "status": {"$in": ["active", "expired"]}},
        sort=[("created_at", -1)]
    )

    if membership:
        update_data = {
            "end_date": new_end_date,
            "modified_at": datetime.now(timezone.utc).isoformat(),
            "modified_by": admin.get("email", admin.get("id")),
        }
        if comment:
            update_data["modification_comment"] = comment

        # Check if new date makes it active or expired
        from datetime import datetime as dt
        try:
            end_dt = dt.strptime(new_end_date, "%Y-%m-%d")
            update_data["status"] = "active" if end_dt >= dt.now() else "expired"
        except ValueError:
            pass

        await db.memberships.update_one(
            {"id": membership["id"]},
            {"$set": update_data}
        )

        # Log the change
        log_entry = {
            "id": str(uuid.uuid4()),
            "type": "membership_change",
            "member_id": member_id,
            "member_name": member.get("name"),
            "gym_id": member.get("gym_id"),
            "previous_end_date": membership.get("end_date"),
            "new_end_date": new_end_date,
            "comment": comment,
            "changed_by": admin.get("email", admin.get("id")),
            "changed_by_role": admin.get("role"),
            "created_at": datetime.now(timezone.utc).isoformat()
        }
        await db.membership_logs.insert_one(log_entry)

        return {"success": True, "message": f"Vencimiento actualizado a {new_end_date}"}
    else:
        raise HTTPException(status_code=404, detail="No se encontro membresia activa para este socio")


@router.get("/members/{member_id}/membership-logs")
async def get_membership_logs(member_id: str, admin: dict = Depends(get_current_admin)):
    """Get membership change history for a member."""
    check_permission(admin, "members_view")

    member = await db.members.find_one({"id": member_id}, {"_id": 0})
    if not member:
        raise HTTPException(status_code=404, detail="Socio no encontrado")

    if admin["role"] != "super_admin" and member.get("gym_id") != admin.get("gym_id"):
        raise HTTPException(status_code=403, detail="No tienes acceso a este socio")

    logs = await db.membership_logs.find(
        {"member_id": member_id},
        {"_id": 0}
    ).sort("created_at", -1).to_list(50)

    return logs



@router.get("/members/{member_id}")
async def get_member(member_id: str, admin: dict = Depends(get_current_admin)):
    member = await db.members.find_one({"id": member_id}, {"_id": 0})
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")
    return member

@router.put("/members/{member_id}")
async def update_member(member_id: str, member_update: MemberUpdate, admin: dict = Depends(get_current_admin)):
    check_permission(admin, "members_edit")
    update_data = {k: v for k, v in member_update.model_dump().items() if v is not None}
    if not update_data:
        raise HTTPException(status_code=400, detail="No data to update")
    update_data["updated_at"] = datetime.now(timezone.utc).isoformat()
    await db.members.update_one({"id": member_id}, {"$set": update_data})
    member = await db.members.find_one({"id": member_id}, {"_id": 0})
    return member

@router.post("/members/{member_id}/approve")
async def approve_member(member_id: str, admin: dict = Depends(get_current_admin)):
    await db.members.update_one({"id": member_id}, {"$set": {"status": "active"}, "$unset": {"suspension_type": "", "suspension_reason": "", "suspended_at": "", "suspended_by": ""}})
    return {"message": "Member approved"}

@router.post("/members/{member_id}/block")
async def block_member(member_id: str, admin: dict = Depends(get_current_admin)):
    await db.members.update_one({"id": member_id}, {"$set": {"status": "blocked"}})
    return {"message": "Member blocked"}

@router.post("/members/{member_id}/suspend")
async def suspend_member(member_id: str, body: dict = {}, admin: dict = Depends(get_current_admin)):
    check_permission(admin, "members_suspend")
    reason = body.get("reason", "Sin motivo especificado") if isinstance(body, dict) else "Sin motivo especificado"
    await db.members.update_one({"id": member_id}, {"$set": {
        "status": "suspended",
        "suspension_type": "manual",
        "suspension_reason": reason,
        "suspended_at": datetime.now(timezone.utc).isoformat(),
        "suspended_by": admin.get("name", admin.get("email", ""))
    }})
    return {"message": "Member suspended", "reason": reason}

@router.delete("/members/{member_id}")
async def delete_member(member_id: str, admin: dict = Depends(get_current_admin)):
    check_permission(admin, "members_delete")
    member = await db.members.find_one({"id": member_id})
    if not member:
        raise HTTPException(status_code=404, detail="Member not found")
    await db.members.delete_one({"id": member_id})
    await db.memberships.delete_many({"member_id": member_id})
    await db.bookings.delete_many({"member_id": member_id})
    await db.guests.delete_many({"member_id": member_id})
    await db.access_logs.delete_many({"member_id": member_id})
    return {"message": "Member deleted"}

@router.post("/members/check-expired-memberships")
async def check_expired_memberships(admin: dict = Depends(get_current_admin)):
    now = datetime.now(timezone.utc).isoformat()
    active_memberships = await db.memberships.find({"status": "active"}, {"_id": 0}).to_list(10000)
    suspended_count = 0
    for m in active_memberships:
        if m.get("end_date") and m["end_date"] < now:
            await db.memberships.update_one({"id": m["id"]}, {"$set": {"status": "expired"}})
            other_active = await db.memberships.find_one({"member_id": m["member_id"], "status": "active"})
            if not other_active:
                await db.members.update_one({"id": m["member_id"]}, {"$set": {
                    "status": "suspended",
                    "suspension_type": "payment",
                    "suspension_reason": "Membresia vencida",
                    "suspended_at": now
                }})
                suspended_count += 1
    return {"message": f"{suspended_count} members suspended due to expired memberships"}

@router.post("/members/cleanup-inactive")
async def cleanup_inactive_members(body: dict = {}, admin: dict = Depends(get_current_admin)):
    """Delete members inactive for 60+ days with no payments and no access"""
    days = body.get("days", 60) if isinstance(body, dict) else 60
    cutoff = (datetime.now(timezone.utc) - timedelta(days=days)).isoformat()
    
    query = {"status": {"$in": ["suspended", "pending", "blocked"]}}
    if admin["role"] != "super_admin":
        query["gym_id"] = admin.get("gym_id")
    
    candidates = await db.members.find(query, {"_id": 0}).to_list(10000)
    deleted_count = 0
    deleted_names = []
    
    for member in candidates:
        member_id = member["id"]
        # Check if created more than X days ago
        created = member.get("created_at", "")
        if created > cutoff:
            continue
        # Check if any recent access
        recent_access = await db.access_logs.find_one({
            "member_id": member_id, "timestamp": {"$gte": cutoff}
        })
        if recent_access:
            continue
        # Check if any paid transaction ever
        paid_tx = await db.payment_transactions.find_one({
            "member_id": member_id, "payment_status": "paid"
        })
        if paid_tx:
            continue
        # Safe to delete
        await db.members.delete_one({"id": member_id})
        await db.memberships.delete_many({"member_id": member_id})
        await db.bookings.delete_many({"member_id": member_id})
        await db.access_logs.delete_many({"member_id": member_id})
        await db.payment_transactions.delete_many({"member_id": member_id})
        await db.email_logs.delete_many({"member_id": member_id})
        deleted_count += 1
        deleted_names.append(member.get("name", member_id))
    
    return {
        "message": f"{deleted_count} socios inactivos eliminados",
        "deleted_count": deleted_count,
        "deleted_members": deleted_names[:50],
        "criteria": f"Inactivos {days}+ dias, sin pagos, sin accesos"
    }


@router.put("/members/{member_id}/rfid")
async def assign_rfid(member_id: str, body: dict, admin: dict = Depends(get_current_admin)):
    """Assign or remove an RFID UID to a member"""
    rfid_uid = body.get("rfid_uid", "").strip().upper().replace(":", "")
    member = await db.members.find_one({"id": member_id}, {"_id": 0})
    if not member:
        raise HTTPException(status_code=404, detail="Socio no encontrado")
    if rfid_uid:
        # Check if RFID is already assigned to another member
        existing = await db.members.find_one({"rfid_uid": rfid_uid, "id": {"$ne": member_id}}, {"_id": 0})
        if existing:
            raise HTTPException(status_code=400, detail=f"Esta tarjeta RFID ya esta asignada a {existing.get('name','otro socio')}")
    await db.members.update_one({"id": member_id}, {"$set": {"rfid_uid": rfid_uid if rfid_uid else None}})
    return {"message": "RFID actualizado" if rfid_uid else "RFID eliminado", "rfid_uid": rfid_uid or None}


@router.get("/members/export/excel")
async def export_members_excel(
    gym_id: Optional[str] = None,
    status: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    include_memberships: bool = False,
    admin: dict = Depends(get_current_admin)
):
    check_permission(admin, "data_export")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    query = {}
    if admin["role"] != "super_admin":
        query["gym_id"] = admin.get("gym_id")
    elif gym_id:
        query["gym_id"] = gym_id
    
    if status and status != "all":
        query["status"] = status
    
    if date_from or date_to:
        date_q = {}
        if date_from:
            date_q["$gte"] = date_from
        if date_to:
            date_q["$lte"] = date_to + "T23:59:59"
        if date_q:
            query["created_at"] = date_q
    
    members = await db.members.find(query, {"_id": 0}).sort("name", 1).to_list(50000)
    
    gym_name = "Todos"
    if gym_id:
        gym = await db.gyms.find_one({"id": gym_id}, {"_id": 0, "name": 1})
        if gym:
            gym_name = gym["name"]
    elif admin["role"] != "super_admin" and admin.get("gym_id"):
        gym = await db.gyms.find_one({"id": admin["gym_id"]}, {"_id": 0, "name": 1})
        if gym:
            gym_name = gym["name"]
    
    membership_map = {}
    if include_memberships and members:
        member_ids = [m["id"] for m in members]
        memberships = await db.memberships.find(
            {"member_id": {"$in": member_ids}, "status": "active"}, {"_id": 0}
        ).to_list(50000)
        plan_ids = list({m["plan_id"] for m in memberships if m.get("plan_id")})
        plans = {}
        if plan_ids:
            for p in await db.plans.find({"id": {"$in": plan_ids}}, {"_id": 0}).to_list(500):
                plans[p["id"]] = p
        for ms in memberships:
            mid = ms["member_id"]
            plan = plans.get(ms.get("plan_id"), {})
            membership_map[mid] = {
                "plan_name": plan.get("name", "-"),
                "start_date": ms.get("start_date", "")[:10],
                "end_date": ms.get("end_date", "")[:10],
                "price": plan.get("price", 0)
            }
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Socios"
    
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )
    
    headers = ["Nombre", "Numero de Socio", "Telefono", "Email", "Estado", "Fecha Registro"]
    if include_memberships:
        headers += ["Plan Activo", "Inicio Plan", "Fin Plan", "Precio Plan"]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    
    status_map = {"active": "Activo", "suspended": "Suspendido", "pending": "Pendiente", "blocked": "Bloqueado"}
    
    for row, member in enumerate(members, 2):
        ws.cell(row=row, column=1, value=member.get("name", "")).border = thin_border
        ws.cell(row=row, column=2, value=member.get("code", "")).border = thin_border
        ws.cell(row=row, column=3, value=member.get("phone", "")).border = thin_border
        ws.cell(row=row, column=4, value=member.get("email", "")).border = thin_border
        ws.cell(row=row, column=5, value=status_map.get(member.get("status", ""), member.get("status", ""))).border = thin_border
        created = member.get("created_at", "")
        ws.cell(row=row, column=6, value=created[:10] if created else "").border = thin_border
        
        if include_memberships:
            ms = membership_map.get(member["id"], {})
            ws.cell(row=row, column=7, value=ms.get("plan_name", "-")).border = thin_border
            ws.cell(row=row, column=8, value=ms.get("start_date", "-")).border = thin_border
            ws.cell(row=row, column=9, value=ms.get("end_date", "-")).border = thin_border
            ws.cell(row=row, column=10, value=ms.get("price", "")).border = thin_border
    
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 35
    ws.column_dimensions['E'].width = 14
    ws.column_dimensions['F'].width = 16
    if include_memberships:
        ws.column_dimensions['G'].width = 22
        ws.column_dimensions['H'].width = 14
        ws.column_dimensions['I'].width = 14
        ws.column_dimensions['J'].width = 14
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    safe_name = gym_name.replace(" ", "_").replace("/", "_")
    filename = f"socios_{safe_name}_{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx"
    
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

@router.put("/members/{member_id}/guest-permission")
async def update_guest_permission(
    member_id: str, 
    can_bring_guests: bool,
    max_guests_per_month: int = 2,
    admin: dict = Depends(get_current_admin)
):
    check_role(admin, ["super_admin", "gym_admin", "gym_manager"])
    await db.members.update_one(
        {"id": member_id},
        {"$set": {"can_bring_guests": can_bring_guests, "max_guests_per_month": max_guests_per_month}}
    )
    return {"message": "Guest permission updated"}


@router.post("/members/import")
async def import_members(
    file: UploadFile = File(...),
    gym_id: str = Form(...),
    keep_codes: bool = Form(True),
    admin: dict = Depends(get_current_admin)
):
    check_role(admin, ["super_admin"])

    gym = await db.gyms.find_one({"id": gym_id})
    if not gym:
        raise HTTPException(status_code=404, detail="Gimnasio no encontrado")

    import openpyxl

    try:
        content = await file.read()
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error al leer el archivo Excel: {str(e)}")

    headers = [cell.value for cell in ws[1]]
    if not headers or len(headers) < 2:
        raise HTTPException(status_code=400, detail="El archivo no tiene headers validos")

    header_map = {}
    for i, h in enumerate(headers):
        if not h:
            continue
        hl = h.lower().strip()
        if 'nombre' in hl and 'completo' in hl:
            header_map['nombre_completo'] = i
        elif 'nombre' in hl and 'completo' not in hl and 'factura' not in hl:
            header_map['nombre'] = i
        elif 'email' in hl and 'factura' not in hl:
            header_map['email'] = i
        elif 'tel' in hl:
            header_map['telefono'] = i
        elif 'dni' in hl or 'documento' in hl:
            header_map['dni'] = i
        elif 'nacimiento' in hl:
            header_map['fecha_nacimiento'] = i
        elif 'alta' in hl:
            header_map['fecha_alta'] = i
        elif 'baja' in hl:
            header_map['fecha_baja'] = i
        elif hl in ['codigo', 'código', 'code']:
            header_map['codigo'] = i
        elif 'socio' in hl and 'num' in hl:
            header_map['numero_socio'] = i
        elif 'cuota' in hl and 'tipo' not in hl:
            header_map['cuota'] = i
        elif 'tipo' in hl and 'cuota' in hl:
            header_map['tipo_cuota'] = i
        elif 'activo' in hl:
            header_map['activo'] = i
        elif 'observ' in hl:
            header_map['observaciones'] = i

    imported = 0
    skipped = 0
    errors = []
    existing_codes = set()

    if keep_codes:
        cursor = db.members.find({"gym_id": gym_id}, {"code": 1, "_id": 0})
        async for doc in cursor:
            if doc.get("code"):
                existing_codes.add(doc["code"])

    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        try:
            nombre_completo = ''
            if 'nombre_completo' in header_map:
                nombre_completo = str(row[header_map['nombre_completo']] or '').strip()
            elif 'nombre' in header_map:
                nombre_completo = str(row[header_map['nombre']] or '').strip()

            if not nombre_completo:
                continue

            email = str(row[header_map['email']] or '').strip() if 'email' in header_map else ''
            phone = str(row[header_map['telefono']] or '').strip() if 'telefono' in header_map else ''
            dni = str(row[header_map['dni']] or '').strip() if 'dni' in header_map else ''

            # Skip if email already exists in this gym
            if email:
                existing = await db.members.find_one({"email": email, "gym_id": gym_id})
                if existing:
                    skipped += 1
                    continue

            # Handle member code
            code = ''
            if keep_codes and 'codigo' in header_map:
                code = str(row[header_map['codigo']] or '').strip()
                if code and code in existing_codes:
                    code = ''  # Code already taken, generate new one

            if not code:
                code = generate_member_code()
                while await db.members.find_one({"code": code}):
                    code = generate_member_code()

            existing_codes.add(code)

            # Determine status
            activo_val = str(row[header_map['activo']] or '').strip().lower() if 'activo' in header_map else 'si'
            status = 'active' if activo_val in ['si', 'sí', 'yes', 'true', '1', 'activo'] else 'suspended'

            cuota = str(row[header_map['cuota']] or '').strip() if 'cuota' in header_map else ''
            tipo_cuota = str(row[header_map['tipo_cuota']] or '').strip() if 'tipo_cuota' in header_map else ''
            observaciones = str(row[header_map['observaciones']] or '').strip() if 'observaciones' in header_map else ''
            fecha_alta = str(row[header_map['fecha_alta']] or '').strip() if 'fecha_alta' in header_map else ''

            member_dict = {
                "id": str(uuid.uuid4()),
                "name": nombre_completo,
                "email": email if email else None,
                "phone": phone if phone else None,
                "dni": dni if dni else None,
                "gym_id": gym_id,
                "code": code,
                "status": status,
                "gender": "prefer_not_to_say",
                "imported": True,
                "import_source": "ismygym",
                "import_data": {
                    "cuota": cuota,
                    "tipo_cuota": tipo_cuota,
                    "fecha_alta_original": fecha_alta,
                    "observaciones": observaciones,
                },
                "created_at": datetime.now(timezone.utc).isoformat()
            }

            if status == 'suspended':
                member_dict["suspension_type"] = "manual"
                member_dict["suspension_reason"] = "Importado como inactivo desde IsMyGym"

            await db.members.insert_one(member_dict)
            member_dict.pop("_id", None)
            imported += 1

        except Exception as e:
            errors.append(f"Fila {row_num}: {str(e)}")
            if len(errors) > 50:
                break

    return {
        "success": True,
        "imported": imported,
        "skipped": skipped,
        "errors": len(errors),
        "error_details": errors[:10],
        "message": f"Importacion completada: {imported} socios importados, {skipped} omitidos (email duplicado), {len(errors)} errores"
    }


@router.post("/members/assign-memberships-bulk")
async def assign_memberships_bulk(data: dict, admin: dict = Depends(get_current_admin)):
    """Assign memberships to imported members using vencimientos data"""
    check_role(admin, ["super_admin"])

    gym_id = data.get("gym_id")
    vencimientos = data.get("vencimientos", {})
    
    if not gym_id or not vencimientos:
        raise HTTPException(status_code=400, detail="gym_id y vencimientos requeridos")

    gym = await db.gyms.find_one({"id": gym_id})
    if not gym:
        raise HTTPException(status_code=404, detail="Gimnasio no encontrado")

    plans = await db.plans.find({"gym_id": gym_id, "active": True}).to_list(100)
    
    assigned = 0
    skipped = 0
    no_plan = 0
    errors = []

    for ismygym_id, info in vencimientos.items():
        try:
            fecha_hasta = info.get("fecha_hasta", "").strip()
            if not fecha_hasta:
                skipped += 1
                continue

            codigo = info.get("codigo", "").strip()
            email = info.get("email", "").strip()
            importe_str = info.get("importe", "0").replace("€", "").replace(",", ".").strip()
            
            try:
                importe = float(importe_str) if importe_str else 0
            except:
                importe = 0

            # Find member by code or email
            member = None
            if codigo:
                member = await db.members.find_one({"code": codigo, "gym_id": gym_id})
            if not member and email:
                member = await db.members.find_one({"email": email, "gym_id": gym_id})
            
            if not member:
                skipped += 1
                continue

            # Check if already has active membership
            existing = await db.memberships.find_one({
                "member_id": member["id"],
                "status": "active"
            })
            if existing:
                skipped += 1
                continue

            # Find matching plan by price
            matched_plan = None
            for p in plans:
                if abs(p.get("price", 0) - importe) < 0.5:
                    matched_plan = p
                    break
            
            if not matched_plan and plans:
                matched_plan = plans[0]

            if not matched_plan:
                no_plan += 1
                continue

            # Parse fecha_hasta (dd/mm/yyyy)
            try:
                parts = fecha_hasta.split("/")
                end_date = f"{parts[2]}-{parts[1]}-{parts[0]}"
            except:
                errors.append(f"Fecha invalida: {fecha_hasta} ({info.get('nombre', '')})")
                continue

            # Calculate start date from end date minus plan duration
            from datetime import datetime as dt
            end_dt = dt.strptime(end_date, "%Y-%m-%d")
            start_dt = end_dt - timedelta(days=matched_plan.get("duration_days", 30))

            membership = {
                "id": str(uuid.uuid4()),
                "member_id": member["id"],
                "plan_id": matched_plan["id"],
                "gym_id": gym_id,
                "start_date": start_dt.strftime("%Y-%m-%d"),
                "end_date": end_date,
                "status": "active" if end_dt >= dt.now() else "expired",
                "payment_status": "paid",
                "imported": True,
                "created_at": datetime.now(timezone.utc).isoformat()
            }

            await db.memberships.insert_one(membership)
            membership.pop("_id", None)
            assigned += 1

        except Exception as e:
            errors.append(f"{info.get('nombre', '')}: {str(e)}")

    return {
        "success": True,
        "assigned": assigned,
        "skipped": skipped,
        "no_plan": no_plan,
        "errors": len(errors),
        "error_details": errors[:10],
        "message": f"Asignacion completada: {assigned} membresias creadas, {skipped} omitidos, {no_plan} sin plan, {len(errors)} errores"
    }
