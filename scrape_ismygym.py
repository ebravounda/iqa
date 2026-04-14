#!/usr/bin/env python3
"""
Script de migración: Extrae socios de IsMyGym y genera Excel
"""
import requests
from bs4 import BeautifulSoup
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import re
import time
import json
import sys

BASE_URL = "https://lafabrika.ismygym.com"
LOGIN_URL = f"{BASE_URL}/admin/socio_listado.php?r=1"
LIST_URL = f"{BASE_URL}/admin/socio_listado.php"
FORM_URL = f"{BASE_URL}/admin/socio_formulario.php"

USERNAME = "edu"
PASSWORD = "Edu@123456"

OUTPUT_FILE = "/app/socios_lafabrika.xlsx"
PROGRESS_FILE = "/app/scraping_progress.json"

def save_progress(data):
    with open(PROGRESS_FILE, 'w') as f:
        json.dump(data, f)

def login(session):
    print("Iniciando sesión en IsMyGym...")
    # Get login page for cookies
    session.get(LOGIN_URL, verify=False)
    
    # POST to correct login endpoint
    login_data = {
        'user': USERNAME,
        'pass': PASSWORD,
        'recordar': 'recordar'
    }
    resp = session.post(f"{BASE_URL}/usuarios/registro.php", data=login_data, verify=False, allow_redirects=True)
    
    if 'Cerrar' in resp.text:
        print("Login exitoso!")
        return True
    
    print("ERROR: No se pudo iniciar sesión")
    return False

def get_member_ids(session):
    print("Recopilando IDs de socios...")
    all_ids = []
    page = 0
    
    while True:
        url = f"{LIST_URL}?invitado=false&pag={page}"
        resp = session.get(url, verify=False)
        soup = BeautifulSoup(resp.text, 'html.parser')
        
        # Find edit links
        edit_links = soup.find_all('a', href=re.compile(r'socio_formulario\.php\?id=\d+'))
        
        page_ids = set()
        for link in edit_links:
            href = link.get('href', '')
            match = re.search(r'id=(\d+)', href)
            if match:
                page_ids.add(match.group(1))
        
        if not page_ids:
            print(f"  Página {page}: sin socios, fin del listado")
            break
        
        all_ids.extend(list(page_ids))
        print(f"  Página {page}: {len(page_ids)} socios encontrados (total: {len(all_ids)})")
        
        # Check if there's a next page
        next_link = soup.find('a', href=re.compile(f'pag={page+1}'))
        if not next_link:
            break
        
        page += 1
        time.sleep(0.5)
    
    # Remove duplicates preserving order
    seen = set()
    unique_ids = []
    for mid in all_ids:
        if mid not in seen:
            seen.add(mid)
            unique_ids.append(mid)
    
    print(f"Total socios únicos: {len(unique_ids)}")
    return unique_ids

def get_member_details(session, member_id):
    url = f"{FORM_URL}?id={member_id}"
    resp = session.get(url, verify=False)
    soup = BeautifulSoup(resp.text, 'html.parser')
    
    def get_field(name):
        inp = soup.find('input', {'name': name})
        if inp:
            return inp.get('value', '').strip()
        sel = soup.find('select', {'name': name})
        if sel:
            opt = sel.find('option', selected=True)
            if opt:
                return opt.text.strip()
        return ''
    
    def get_textarea(name):
        ta = soup.find('textarea', {'name': name})
        if ta:
            return ta.text.strip()
        return ''
    
    member = {
        'id_ismygym': member_id,
        'nombre': get_field('Nombre'),
        'apellidos': get_field('Apellidos'),
        'email': get_field('Email'),
        'telefono': get_field('Telefono'),
        'dni': get_field('Dni'),
        'fecha_nacimiento': get_field('FechaNacimiento'),
        'fecha_alta': get_field('FechaAlta'),
        'fecha_baja': get_field('FechaBaja'),
        'numero_socio': get_field('NumeroSocio'),
        'codigo': get_field('CodigoReserva'),
        'importe_cuota': get_field('ImporteCuota'),
        'tipo_cuota': get_field('TipoCuota'),
        'forma_pago': get_field('FormaPago'),
        'activo': 'Si' if soup.find('input', {'name': 'Activo', 'checked': True}) else 'No',
        'observaciones': get_textarea('Observaciones'),
    }
    
    # Check active status from checkbox
    activo_input = soup.find('input', {'name': 'Activo'})
    if activo_input:
        if activo_input.get('checked') is not None:
            member['activo'] = 'Si'
        else:
            member['activo'] = 'No'
    
    return member

def create_excel(members):
    print(f"\nGenerando Excel con {len(members)} socios...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Socios La Fabrika"
    
    # Headers
    headers = [
        'Nombre Completo', 'Email', 'Teléfono', 'DNI',
        'Fecha Nacimiento', 'Fecha Alta', 'Fecha Baja',
        'Código', 'Nº Socio', 'Cuota', 'Tipo Cuota',
        'Forma Pago', 'Activo', 'Observaciones'
    ]
    
    # Style
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = border
    
    # Data
    for row, m in enumerate(members, 2):
        nombre_completo = f"{m.get('nombre', '')} {m.get('apellidos', '')}".strip()
        values = [
            nombre_completo,
            m.get('email', ''),
            m.get('telefono', ''),
            m.get('dni', ''),
            m.get('fecha_nacimiento', ''),
            m.get('fecha_alta', ''),
            m.get('fecha_baja', ''),
            m.get('codigo', ''),
            m.get('numero_socio', ''),
            m.get('importe_cuota', ''),
            m.get('tipo_cuota', ''),
            m.get('forma_pago', ''),
            m.get('activo', ''),
            m.get('observaciones', ''),
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = border
            if col <= 3:
                cell.alignment = Alignment(horizontal='left')
            else:
                cell.alignment = Alignment(horizontal='center')
    
    # Auto-width
    for col in ws.columns:
        max_length = 0
        for cell in col:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = min(len(str(cell.value)), 40)
            except:
                pass
        ws.column_dimensions[col[0].column_letter].width = max_length + 4
    
    # Freeze header
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    
    wb.save(OUTPUT_FILE)
    print(f"Excel guardado en: {OUTPUT_FILE}")

def main():
    import urllib3
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
    
    session = requests.Session()
    session.headers.update({
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/120.0.0.0'
    })
    
    if not login(session):
        sys.exit(1)
    
    member_ids = get_member_ids(session)
    
    if not member_ids:
        print("No se encontraron socios")
        sys.exit(1)
    
    members = []
    total = len(member_ids)
    errors = 0
    
    for i, mid in enumerate(member_ids):
        try:
            member = get_member_details(session, mid)
            members.append(member)
            
            if (i + 1) % 10 == 0 or i == 0:
                print(f"  Progreso: {i+1}/{total} ({((i+1)/total*100):.1f}%) - Último: {member['nombre']} {member['apellidos']}")
                save_progress({
                    'total': total,
                    'processed': i + 1,
                    'errors': errors,
                    'last_member': f"{member['nombre']} {member['apellidos']}"
                })
            
            time.sleep(0.3)
            
        except Exception as e:
            errors += 1
            print(f"  ERROR en socio {mid}: {str(e)}")
            if errors > 20:
                print("Demasiados errores, deteniendo...")
                break
    
    if members:
        create_excel(members)
        save_progress({
            'total': total,
            'processed': len(members),
            'errors': errors,
            'status': 'COMPLETADO',
            'output_file': OUTPUT_FILE
        })
        print(f"\n{'='*50}")
        print(f"MIGRACIÓN COMPLETADA")
        print(f"Socios extraídos: {len(members)}")
        print(f"Errores: {errors}")
        print(f"Archivo: {OUTPUT_FILE}")
        print(f"{'='*50}")
    else:
        print("No se pudieron extraer datos de ningún socio")

if __name__ == '__main__':
    main()
